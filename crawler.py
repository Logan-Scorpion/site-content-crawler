#!/usr/bin/env python3
"""
Site content crawler.

Crawls every reachable HTML page on a website (same domain only) and catalogs
each page's content into an Excel spreadsheet:

  * Page URL
  * Page title (<title>)
  * Meta description
  * Main heading (H1)
  * Full body copy, including subheadings (marked as [H2], [H3], ...)
  * Image alt text

Usage:
    python crawler.py https://www.example.com
    python crawler.py https://www.example.com -o example.xlsx --max-pages 1000
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import time
import xml.etree.ElementTree as ET
from collections import deque
from dataclasses import dataclass, field
from urllib import robotparser
from urllib.parse import urldefrag, urljoin, urlparse, urlunparse

import requests
from bs4 import BeautifulSoup, Comment, NavigableString, Tag

DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (compatible; SiteContentCrawler/1.0; +https://github.com/)"
)

# File extensions that are never HTML pages, so we don't bother requesting them.
SKIP_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".ico", ".bmp", ".tif",
    ".tiff", ".avif", ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx",
    ".zip", ".rar", ".gz", ".tar", ".7z", ".mp3", ".mp4", ".mov", ".avi",
    ".wmv", ".webm", ".wav", ".ogg", ".css", ".js", ".json", ".xml", ".txt",
    ".woff", ".woff2", ".ttf", ".eot", ".otf", ".exe", ".dmg", ".rss", ".atom",
}

# Query parameters that only track visits and would otherwise create duplicates.
TRACKING_PARAMS = re.compile(
    r"^(utm_[a-z]+|gclid|fbclid|msclkid|mc_cid|mc_eid|_ga|_gl|ref|yclid)$", re.I
)

# Elements whose text is never visible page copy.
NON_CONTENT_TAGS = [
    "script", "style", "noscript", "template", "svg", "iframe", "canvas",
    "object", "embed", "head", "meta", "link",
]

# Site-wide chrome (navigation, header, footer, sidebars, cookie banners) that
# repeats on every page. Stripped from body copy unless --keep-boilerplate.
BOILERPLATE_SELECTORS = [
    "nav", "footer", "aside",
    "[role=navigation]", "[role=contentinfo]", "[role=banner]",
    "[role=complementary]", "[role=search]", "[aria-hidden=true]",
    "#cookie-banner", ".cookie-banner", "#onetrust-consent-sdk",
    ".skip-link", ".skip-to-content", ".screen-reader-text", ".sr-only",
    ".visually-hidden",
]

BLOCK_TAGS = {
    "p", "div", "section", "article", "main", "header", "footer", "li", "ul",
    "ol", "dl", "dt", "dd", "table", "tr", "td", "th", "thead", "tbody",
    "blockquote", "pre", "figure", "figcaption", "address", "form",
    "fieldset", "legend", "label", "details", "summary", "hr", "br", "option",
}
HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}

EXCEL_CELL_LIMIT = 32767
ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


@dataclass
class PageRecord:
    url: str
    status: int | str = ""
    title: str = ""
    meta_description: str = ""
    h1: str = ""
    body_copy: str = ""
    word_count: int = 0
    images: list[tuple[str, str]] = field(default_factory=list)  # (src, alt)
    error: str = ""

    @property
    def alt_text_summary(self) -> str:
        alts = [alt for _, alt in self.images if alt]
        return "\n".join(alts)

    @property
    def missing_alt_count(self) -> int:
        return sum(1 for _, alt in self.images if not alt)


# --------------------------------------------------------------------------- #
# URL helpers
# --------------------------------------------------------------------------- #

def bare_host(netloc: str) -> str:
    host = netloc.lower().split("@")[-1].split(":")[0]
    return host[4:] if host.startswith("www.") else host


def normalize_url(url: str) -> str:
    """Canonical form used for de-duplication."""
    url, _ = urldefrag(url.strip())
    parts = urlparse(url)
    scheme = parts.scheme.lower()
    netloc = parts.netloc.lower()
    if (scheme == "http" and netloc.endswith(":80")) or (
        scheme == "https" and netloc.endswith(":443")
    ):
        netloc = netloc.rsplit(":", 1)[0]
    path = parts.path or "/"
    path = re.sub(r"/{2,}", "/", path)
    query = "&".join(
        sorted(
            q for q in parts.query.split("&")
            if q and not TRACKING_PARAMS.match(q.split("=", 1)[0])
        )
    )
    return urlunparse((scheme, netloc, path, "", query, ""))


class Scope:
    """Decides which URLs belong to the site being crawled."""

    def __init__(self, start_url: str, include_subdomains: bool,
                 path_prefix: str | None):
        self.root = bare_host(urlparse(start_url).netloc)
        self.include_subdomains = include_subdomains
        self.path_prefix = path_prefix

    def contains(self, url: str) -> bool:
        parts = urlparse(url)
        if parts.scheme not in ("http", "https"):
            return False
        host = bare_host(parts.netloc)
        if host != self.root and not (
            self.include_subdomains and host.endswith("." + self.root)
        ):
            return False
        if self.path_prefix and not parts.path.startswith(self.path_prefix):
            return False
        last_segment = parts.path.rsplit("/", 1)[-1]
        ext = "." + last_segment.rsplit(".", 1)[-1].lower() if "." in last_segment else ""
        return ext not in SKIP_EXTENSIONS


# --------------------------------------------------------------------------- #
# Content extraction
# --------------------------------------------------------------------------- #

def clean(text: str | None) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def find_content_root(soup: BeautifulSoup) -> Tag:
    """Prefer the page's <main> region; fall back to <body>."""
    for selector in ("main", "[role=main]", "#main-content", "#main", "#content"):
        node = soup.select_one(selector)
        if node and clean(node.get_text()):
            return node
    return soup.body or soup


def strip_boilerplate(root: Tag) -> None:
    for selector in BOILERPLATE_SELECTORS:
        for node in root.select(selector):
            node.decompose()
    # A site-wide <header> lives outside <main>/<article>; headers inside them
    # often hold the page's H1, so keep those.
    for header in root.find_all("header"):
        if not header.find_parent(["main", "article"]) and not header.find("h1"):
            header.decompose()


def extract_body_copy(root: Tag) -> str:
    """Flatten the content tree to readable text, one block per line,
    with subheadings marked like "[H2] Our Services"."""
    lines: list[str] = []
    buffer: list[str] = []

    def flush():
        text = clean("".join(buffer))
        if text:
            lines.append(text)
        buffer.clear()

    def walk(node):
        for child in node.children:
            if isinstance(child, Comment):
                continue
            if isinstance(child, NavigableString):
                buffer.append(str(child))
                continue
            if not isinstance(child, Tag):
                continue
            name = child.name
            if name in HEADING_TAGS:
                flush()
                text = clean(child.get_text(" "))
                if text:
                    lines.append(f"[{name.upper()}] {text}")
            elif name in BLOCK_TAGS:
                flush()
                walk(child)
                flush()
            else:
                walk(child)

    walk(root)
    flush()

    # Collapse consecutive duplicate lines (common with responsive markup).
    deduped = [line for i, line in enumerate(lines) if i == 0 or line != lines[i - 1]]
    return "\n".join(deduped)


def extract_page(url: str, html: str, keep_boilerplate: bool,
                 all_images: bool) -> tuple[PageRecord, list[str]]:
    soup = BeautifulSoup(html, "lxml")
    record = PageRecord(url=url)

    if soup.title:
        record.title = clean(soup.title.get_text())

    meta = soup.find("meta", attrs={"name": re.compile(r"^description$", re.I)})
    if meta is None:
        meta = soup.find("meta", attrs={"property": "og:description"})
    if meta is not None:
        record.meta_description = clean(meta.get("content"))

    # Links are collected from the full page (nav included) before anything is
    # stripped, so the crawler still discovers every page.
    base = soup.find("base", href=True)
    base_url = urljoin(url, base["href"]) if base else url
    links = []
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if href and not href.lower().startswith(("mailto:", "tel:", "javascript:", "data:", "#")):
            links.append(urljoin(base_url, href))

    for tag in soup.find_all(NON_CONTENT_TAGS):
        tag.decompose()

    record.h1 = " | ".join(
        t for t in (clean(h.get_text(" ")) for h in soup.find_all("h1")) if t
    )

    root = find_content_root(soup)
    # Collect images before boilerplate is stripped when cataloging them all.
    all_imgs = (soup.body or soup).find_all("img") if all_images else None
    if not keep_boilerplate:
        strip_boilerplate(root)

    for img in all_imgs if all_imgs is not None else root.find_all("img"):
        src = (
            img.get("src") or img.get("data-src") or img.get("data-lazy-src")
            or (img.get("srcset") or "").split(" ")[0]
        )
        if src and not src.startswith("data:"):
            src = urljoin(base_url, src)
        record.images.append((src or "", clean(img.get("alt"))))

    record.body_copy = extract_body_copy(root)
    record.word_count = len(re.findall(r"\b\w+\b", re.sub(r"^\[H\d\] ", "", record.body_copy, flags=re.M)))
    return record, links


# --------------------------------------------------------------------------- #
# Crawler
# --------------------------------------------------------------------------- #

class Crawler:
    def __init__(self, start_url: str, *, max_pages: int, delay: float,
                 timeout: float, user_agent: str, respect_robots: bool,
                 use_sitemap: bool, include_subdomains: bool,
                 path_prefix: str | None, keep_boilerplate: bool,
                 all_images: bool, verbose: bool = True):
        if not urlparse(start_url).scheme:
            start_url = "https://" + start_url
        self.start_url = normalize_url(start_url)
        self.max_pages = max_pages
        self.delay = delay
        self.timeout = timeout
        self.keep_boilerplate = keep_boilerplate
        self.all_images = all_images
        self.verbose = verbose
        self.scope = Scope(self.start_url, include_subdomains, path_prefix)
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": user_agent,
            "Accept": "text/html,application/xhtml+xml;q=0.9,*/*;q=0.8",
        })
        self.robots = self._load_robots() if respect_robots else None
        self.use_sitemap = use_sitemap
        self.user_agent = user_agent

    def log(self, msg: str) -> None:
        if self.verbose:
            print(msg, file=sys.stderr, flush=True)

    def _load_robots(self):
        parts = urlparse(self.start_url)
        rp = robotparser.RobotFileParser()
        try:
            resp = self.session.get(
                f"{parts.scheme}://{parts.netloc}/robots.txt", timeout=self.timeout
            )
            rp.parse(resp.text.splitlines() if resp.ok else [])
        except requests.RequestException:
            rp.parse([])
        return rp

    def allowed(self, url: str) -> bool:
        return self.robots is None or self.robots.can_fetch(self.user_agent, url)

    def sitemap_urls(self) -> list[str]:
        parts = urlparse(self.start_url)
        root = f"{parts.scheme}://{parts.netloc}"
        candidates = list(self.robots.site_maps() or []) if self.robots else []
        candidates.append(root + "/sitemap.xml")
        seen, found = set(), []
        queue = deque(candidates)
        while queue and len(seen) < 200:
            sm = queue.popleft()
            if sm in seen:
                continue
            seen.add(sm)
            try:
                resp = self.session.get(sm, timeout=self.timeout)
                if not resp.ok:
                    continue
                tree = ET.fromstring(resp.content)
            except (requests.RequestException, ET.ParseError):
                continue
            for loc in tree.iter():
                if not loc.tag.endswith("loc") or not loc.text:
                    continue
                loc_url = loc.text.strip()
                if loc_url.lower().endswith((".xml", ".xml.gz")):
                    queue.append(loc_url)
                else:
                    found.append(loc_url)
        if found:
            self.log(f"Found {len(found)} URLs in sitemap(s)")
        return found

    def crawl(self) -> list[PageRecord]:
        queue: deque[str] = deque([self.start_url])
        queued: set[str] = {self.start_url}
        done: set[str] = set()
        records: list[PageRecord] = []

        if self.use_sitemap:
            for u in self.sitemap_urls():
                n = normalize_url(u)
                if n not in queued and self.scope.contains(n):
                    queued.add(n)
                    queue.append(n)

        while queue and len(records) < self.max_pages:
            url = queue.popleft()
            if url in done:
                continue
            if not self.allowed(url):
                self.log(f"  skip (robots.txt): {url}")
                continue

            try:
                resp = self.session.get(url, timeout=self.timeout, allow_redirects=True)
            except requests.RequestException as exc:
                records.append(PageRecord(url=url, status="error", error=str(exc)[:500]))
                self.log(f"[{len(records)}] ERROR {url}: {exc}")
                continue
            finally:
                if self.delay:
                    time.sleep(self.delay)

            final_url = normalize_url(resp.url)
            if final_url in done:
                continue  # redirected to a page we already have
            if final_url != url and not self.scope.contains(final_url):
                continue  # redirected off-site
            done.add(url)
            done.add(final_url)

            ctype = resp.headers.get("Content-Type", "")
            if resp.ok and "html" not in ctype.lower():
                continue  # PDFs, images, etc. served from extensionless URLs

            if not resp.ok:
                records.append(PageRecord(url=final_url, status=resp.status_code,
                                          error=resp.reason or ""))
                self.log(f"[{len(records)}] {resp.status_code} {final_url}")
                continue

            if resp.encoding is None or resp.encoding.lower() == "iso-8859-1":
                resp.encoding = resp.apparent_encoding
            record, links = extract_page(final_url, resp.text,
                                         self.keep_boilerplate, self.all_images)
            record.status = resp.status_code
            records.append(record)
            self.log(f"[{len(records)}] {resp.status_code} {final_url}")

            # Respect <meta name="robots" content="nofollow">.
            if re.search(r'<meta[^>]+name=["\']robots["\'][^>]+nofollow', resp.text[:20000], re.I):
                continue
            for link in links:
                n = normalize_url(link)
                if n not in queued and self.scope.contains(n):
                    queued.add(n)
                    queue.append(n)

        if queue:
            self.log(f"Stopped at --max-pages={self.max_pages}; "
                     f"{len(queue)} URLs left unvisited.")
        return records


# --------------------------------------------------------------------------- #
# Output
# --------------------------------------------------------------------------- #

PAGE_COLUMNS = [
    ("Page URL", 50, lambda r: r.url),
    ("Status", 8, lambda r: r.status),
    ("Page Title", 40, lambda r: r.title),
    ("Title Length", 8, lambda r: len(r.title)),
    ("Meta Description", 50, lambda r: r.meta_description),
    ("Meta Description Length", 10, lambda r: len(r.meta_description)),
    ("H1", 40, lambda r: r.h1),
    ("Body Copy", 100, lambda r: r.body_copy),
    ("Word Count", 9, lambda r: r.word_count),
    ("Image Alt Text", 50, lambda r: r.alt_text_summary),
    ("Images", 8, lambda r: len(r.images)),
    ("Images Missing Alt", 10, lambda r: r.missing_alt_count),
    ("Error", 30, lambda r: r.error),
]


def excel_safe(value):
    if not isinstance(value, str):
        return value
    value = ILLEGAL_XML_CHARS.sub("", value)
    if value.startswith(("=", "+", "-", "@")):
        value = "'" + value  # prevent formula injection from page text
    if len(value) > EXCEL_CELL_LIMIT:
        note = "\n...[truncated: exceeds Excel cell limit]"
        value = value[: EXCEL_CELL_LIMIT - len(note)] + note
    return value


def write_xlsx(records: list[PageRecord], path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    wrap_top = Alignment(wrap_text=True, vertical="top")

    def setup(ws, columns):
        ws.append([name for name, *_ in columns])
        for i, (_, width, *_) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=i)
            cell.font, cell.fill = header_font, header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    pages = wb.active
    pages.title = "Pages"
    setup(pages, PAGE_COLUMNS)
    for r in records:
        pages.append([excel_safe(getter(r)) for _, _, getter in PAGE_COLUMNS])
    for row in pages.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_top
        pages.row_dimensions[row[0].row].height = 90

    images = wb.create_sheet("Images")
    image_columns = [("Page URL", 50), ("Image URL", 60), ("Alt Text", 60)]
    setup(images, image_columns)
    for r in records:
        for src, alt in r.images:
            images.append([excel_safe(r.url), excel_safe(src),
                           excel_safe(alt) if alt else "(missing)"])

    wb.save(path)


def write_csv(records: list[PageRecord], path: str) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow([name for name, *_ in PAGE_COLUMNS])
        for r in records:
            writer.writerow([getter(r) for _, _, getter in PAGE_COLUMNS])


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def default_output_name(url: str) -> str:
    host = bare_host(urlparse(url if "//" in url else "https://" + url).netloc)
    return f"{host or 'site'}-content.xlsx"


def parse_args(argv=None):
    p = argparse.ArgumentParser(
        description="Crawl a website and catalog each page's content into a spreadsheet."
    )
    p.add_argument("url", help="Start URL, e.g. https://www.example.com")
    p.add_argument("-o", "--output",
                   help="Output file (.xlsx or .csv). Default: <domain>-content.xlsx")
    p.add_argument("--max-pages", type=int, default=500,
                   help="Maximum number of pages to crawl (default: 500)")
    p.add_argument("--delay", type=float, default=0.5,
                   help="Seconds to wait between requests (default: 0.5)")
    p.add_argument("--timeout", type=float, default=20, help="Request timeout in seconds")
    p.add_argument("--user-agent", default=DEFAULT_USER_AGENT)
    p.add_argument("--ignore-robots", action="store_true",
                   help="Do not obey robots.txt")
    p.add_argument("--no-sitemap", action="store_true",
                   help="Do not seed the crawl from sitemap.xml")
    p.add_argument("--include-subdomains", action="store_true",
                   help="Also crawl subdomains (e.g. blog.example.com)")
    p.add_argument("--path-prefix",
                   help="Only crawl URLs whose path starts with this (e.g. /blog/)")
    p.add_argument("--keep-boilerplate", action="store_true",
                   help="Include nav/header/footer/sidebar text in body copy")
    p.add_argument("--all-images", action="store_true",
                   help="Catalog every image on the page, including logos/icons "
                        "in the header and footer")
    p.add_argument("-q", "--quiet", action="store_true", help="Suppress progress output")
    return p.parse_args(argv)


def main(argv=None) -> int:
    args = parse_args(argv)
    output = args.output or default_output_name(args.url)

    crawler = Crawler(
        args.url,
        max_pages=args.max_pages,
        delay=args.delay,
        timeout=args.timeout,
        user_agent=args.user_agent,
        respect_robots=not args.ignore_robots,
        use_sitemap=not args.no_sitemap,
        include_subdomains=args.include_subdomains,
        path_prefix=args.path_prefix,
        keep_boilerplate=args.keep_boilerplate,
        all_images=args.all_images,
        verbose=not args.quiet,
    )
    records = crawler.crawl()

    if output.lower().endswith(".csv"):
        write_csv(records, output)
    else:
        write_xlsx(records, output)
    print(f"Saved {len(records)} pages to {output}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
