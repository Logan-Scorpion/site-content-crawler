"""End-to-end test: serve a tiny site locally, crawl it, check the spreadsheet."""
import functools
import http.server
import sys
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import crawler  # noqa: E402

NAV = '<header><a href="/"><img src="/logo.png" alt="Acme logo"></a><nav><a href="/about/">About</a><a href="/services">Services</a></nav></header>'
FOOT = '<footer>Copyright Acme <a href="/privacy">Privacy</a></footer>'

PAGES = {
    "index.html": f"""<html><head><title>Acme Home</title>
<meta name="description" content="Acme does great things."></head><body>{NAV}
<main><h1>Welcome to Acme</h1><p>We build <strong>rockets</strong>.</p>
<h2>Why us</h2><ul><li>Fast</li><li>Cheap</li></ul>
<img src="/hero.jpg" alt="A rocket launching"><img src="/deco.jpg">
<a href="/contact#form">Contact</a><a href="https://other.com/">Ext</a>
<a href="/brochure.pdf">PDF</a></main>{FOOT}<script>var x=1;</script></body></html>""",
    "about/index.html": f"<html><head><title>About</title></head><body>{NAV}<main><h1>About Acme</h1><p>Since 1950.</p></main>{FOOT}</body></html>",
    "services": f"<html><head><title>Services</title></head><body>{NAV}<h1>Services</h1><h3>Launch</h3><p>To orbit.</p>{FOOT}</body></html>",
    "contact": f"<html><head><title>Contact</title></head><body>{NAV}<main><h1>Contact</h1></main>{FOOT}</body></html>",
    "privacy": f"<html><head><title>Privacy</title></head><body>{NAV}<main><h1>Privacy</h1></main>{FOOT}</body></html>",
    "robots.txt": "User-agent: *\nDisallow: /privacy\n",
}


@pytest.fixture()
def site(tmp_path):
    for name, body in PAGES.items():
        f = tmp_path / name
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(body)

    class Handler(http.server.SimpleHTTPRequestHandler):
        def guess_type(self, path):
            return "text/plain" if path.endswith(".txt") else "text/html"

        def log_message(self, *a):
            pass

    server = http.server.ThreadingHTTPServer(
        ("127.0.0.1", 0), functools.partial(Handler, directory=str(tmp_path)))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}/"
    server.shutdown()


def test_crawl_to_xlsx(site, tmp_path):
    out = tmp_path / "out.xlsx"
    assert crawler.main([site, "-o", str(out), "--delay", "0", "-q"]) == 0

    wb = load_workbook(out)
    rows = list(wb["Pages"].iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    pages = {r[0].replace(site.rstrip("/"), ""): dict(zip(header, r)) for r in data}

    assert set(pages) == {"/", "/about/", "/services", "/contact"}  # /privacy blocked by robots

    home = pages["/"]
    assert home["Page Title"] == "Acme Home"
    assert home["Meta Description"] == "Acme does great things."
    assert home["H1"] == "Welcome to Acme"
    body = home["Body Copy"]
    assert "[H1] Welcome to Acme" in body
    assert "We build rockets." in body
    assert "[H2] Why us" in body and "Fast" in body
    assert "Copyright" not in body and "var x" not in body and "About" not in body
    assert home["Image Alt Text"] == "A rocket launching"
    assert home["Images Missing Alt"] == 1

    assert "[H3] Launch" in pages["/services"]["Body Copy"]
    assert "Copyright" not in pages["/services"]["Body Copy"]

    imgs = list(wb["Images"].iter_rows(values_only=True))[1:]
    assert (home["Page URL"], site + "deco.jpg", "(missing)") in imgs


def test_all_images_includes_logo(site, tmp_path):
    out = tmp_path / "out.csv"
    crawler.main([site, "-o", str(out), "--delay", "0", "-q", "--all-images", "--max-pages", "1"])
    text = out.read_text(encoding="utf-8-sig")
    assert "Acme logo" in text and "A rocket launching" in text


def test_normalize_url():
    n = crawler.normalize_url
    assert n("HTTPS://Example.com:443/a#x") == "https://example.com/a"
    assert n("https://example.com/?utm_source=x&b=2&a=1") == "https://example.com/?a=1&b=2"
